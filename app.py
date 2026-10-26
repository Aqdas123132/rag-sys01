import hashlib
import html
import io
import json
import os
import tempfile
import threading
import time
import uuid
from datetime import datetime
from pathlib import Path

import streamlit as st
from langchain_community.document_loaders import PyPDFLoader

from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openai import (
    APIConnectionError,
    APIStatusError,
    AuthenticationError,
    OpenAI,
    PermissionDeniedError,
    RateLimitError,
)

# =============================================================================
# Page setup
# =============================================================================
st.set_page_config(
    page_title="Document RAG Agent",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

# =============================================================================
# Theme: royal blue, red, white, with a yellow accent  (+ animations)
# =============================================================================
BASE_CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Manrope:wght@400;500;600;700;800&display=swap');

:root {
  --blue: #2150C8;
  --blue-dark: #16399A;
  --red: #D62839;
  --red-dark: #B01E2E;
  --yellow: #FFC72C;
  --ink: #0F1B3D;
  --muted: #5B6787;
  --line: #DCE3F5;
  --soft: #F3F6FF;
  --ease: cubic-bezier(0.22, 1, 0.36, 1);
}

/* ---------- Keyframes ---------- */
@keyframes fadeUp   { from { opacity: 0; transform: translateY(18px); } to { opacity: 1; transform: none; } }
@keyframes fadeIn   { from { opacity: 0; } to { opacity: 1; } }
@keyframes slideLeft { from { opacity: 0; transform: translateX(-26px); } to { opacity: 1; transform: none; } }
@keyframes slideRight { from { opacity: 0; transform: translateX(26px); } to { opacity: 1; transform: none; } }
@keyframes popIn    { 0% { opacity: 0; transform: scale(0.92); } 100% { opacity: 1; transform: scale(1); } }
@keyframes floaty   { 0%,100% { transform: translateY(0); } 50% { transform: translateY(-10px); } }
@keyframes drift    { 0%,100% { transform: translate(0,0) scale(1); } 50% { transform: translate(24px,-28px) scale(1.12); } }
@keyframes gradientShift { 0% { background-position: 0% 50%; } 50% { background-position: 100% 50%; } 100% { background-position: 0% 50%; } }
@keyframes pulseDot { 0% { box-shadow: 0 0 0 0 rgba(255,199,44,0.7); } 70% { box-shadow: 0 0 0 9px rgba(255,199,44,0); } 100% { box-shadow: 0 0 0 0 rgba(255,199,44,0); } }
@keyframes shine    { 0% { transform: translateX(-120%) skewX(-20deg); } 60%,100% { transform: translateX(260%) skewX(-20deg); } }
@keyframes shake    { 0%,100% { transform: translateX(0); } 20% { transform: translateX(-8px); } 40% { transform: translateX(8px); } 60% { transform: translateX(-5px); } 80% { transform: translateX(5px); } }
@keyframes barGrow  { from { transform: scaleX(0); } to { transform: scaleX(1); } }
@keyframes bounce   { 0%,80%,100% { transform: translateY(0); opacity: .5; } 40% { transform: translateY(-8px); opacity: 1; } }
@keyframes wiggle   { 0%,100% { transform: rotate(0); } 25% { transform: rotate(-8deg); } 75% { transform: rotate(8deg); } }

.stApp { font-family: 'Manrope', 'Segoe UI', sans-serif; color: var(--ink); background: #FFFFFF; }
.stApp h1, .stApp h2, .stApp h3, .stApp h4,
.stApp input, .stApp textarea, .stApp button, .stApp label { font-family: 'Manrope', 'Segoe UI', sans-serif; }

[data-testid="stDecoration"], [data-testid="stAppDeployButton"], #MainMenu, footer { display: none !important; }
[data-testid="stHeader"] { background: transparent; }
[data-testid="stExpandSidebarButton"], [data-testid="stSidebarCollapsedControl"],
[data-testid="collapsedControl"], [data-testid="stSidebarCollapseButton"] { visibility: visible !important; }
[data-testid="stExpandSidebarButton"] button, [data-testid="stSidebarCollapsedControl"] button,
[data-testid="stSidebarCollapseButton"] button { color: var(--blue) !important; }
.block-container { padding-top: 2rem; padding-bottom: 3rem; max-width: 1180px; }

/* ---------- Brand ---------- */
.brand { display: flex; align-items: center; gap: 12px; }
.brand-name { font-weight: 800; font-size: 1.2rem; letter-spacing: -0.01em; }
.brand svg { animation: floaty 4s ease-in-out infinite; }

/* ---------- Buttons ---------- */
button[kind="primary"],
[data-testid="stBaseButton-primary"],
[data-testid="stBaseButton-primaryFormSubmit"] {
  position: relative; overflow: hidden;
  background: var(--red) !important; color: #fff !important; border: 0 !important;
  border-radius: 10px !important; font-weight: 700 !important; padding: 0.65rem 1.2rem !important;
  box-shadow: 0 2px 0 var(--red-dark) !important;
  transition: transform .2s var(--ease), box-shadow .2s var(--ease), background .2s !important;
}
button[kind="primary"]::after,
[data-testid="stBaseButton-primaryFormSubmit"]::after {
  content: ""; position: absolute; top: 0; left: 0; width: 40%; height: 100%;
  background: linear-gradient(90deg, transparent, rgba(255,255,255,0.35), transparent);
  animation: shine 3.2s ease-in-out infinite; pointer-events: none;
}
button[kind="primary"]:hover,
[data-testid="stBaseButton-primary"]:hover,
[data-testid="stBaseButton-primaryFormSubmit"]:hover {
  background: var(--red-dark) !important; transform: translateY(-2px);
  box-shadow: 0 8px 18px rgba(214,40,57,0.35) !important;
}
button[kind="primary"]:active, [data-testid="stBaseButton-primaryFormSubmit"]:active { transform: translateY(0) scale(0.98); }
button[kind="primary"]:focus-visible,
[data-testid="stBaseButton-primaryFormSubmit"]:focus-visible { outline: 3px solid var(--yellow) !important; outline-offset: 2px; }

[data-testid="stBaseButton-secondary"], [data-testid="stDownloadButton"] button {
  background: #fff !important; color: var(--blue) !important; border: 1.5px solid var(--blue) !important;
  border-radius: 10px !important; font-weight: 700 !important;
  transition: transform .2s var(--ease), background .2s, color .2s, box-shadow .2s !important;
}
[data-testid="stBaseButton-secondary"]:hover, [data-testid="stDownloadButton"] button:hover {
  background: var(--blue) !important; color: #fff !important; transform: translateY(-2px);
  box-shadow: 0 8px 16px rgba(33,80,200,0.25) !important;
}

/* ---------- Inputs ---------- */
[data-testid="stTextInput"] input {
  border-radius: 10px; border: 1.5px solid var(--line); padding: 0.7rem 0.9rem; background: #fff;
  transition: border-color .2s, box-shadow .2s, transform .2s var(--ease);
}
[data-testid="stTextInput"] input:focus { border-color: var(--blue); box-shadow: 0 0 0 3px rgba(33,80,200,0.15); }
[data-testid="stTextInput"] label p, [data-testid="stSelectbox"] label p,
[data-testid="stSlider"] label p { font-weight: 700; color: var(--ink); }
[data-testid="stSelectbox"] div[data-baseweb="select"] > div { border-radius: 10px; border-color: var(--line); }

/* ---------- Sidebar ---------- */
[data-testid="stSidebar"] { background: var(--soft); border-right: 1px solid var(--line); animation: slideLeft .5s var(--ease) both; }
[data-testid="stSidebar"] .block-container { padding-top: 1.2rem; }
.side-brand { display: flex; align-items: center; gap: 10px; margin: 0 0 14px 0; }
.side-brand svg { width: 32px; height: 32px; }
.side-brand b { font-weight: 800; font-size: 1.02rem; color: var(--ink); }
.side-title {
  font-weight: 800; font-size: 0.95rem; color: var(--blue); margin: 0.4rem 0 0.8rem 0;
  padding-bottom: 6px; border-bottom: 3px solid var(--red); display: inline-block;
}
.key-box {
  background: #fff; border: 1px solid var(--line); border-left: 4px solid var(--yellow);
  border-radius: 10px; padding: 10px 12px; margin-bottom: 10px; font-size: 0.85rem;
  animation: popIn .5s var(--ease) both;
}
.key-box .k-label { color: var(--muted); font-size: 0.75rem; font-weight: 600; }
.key-box .k-value { font-family: ui-monospace, Menlo, Consolas, monospace; font-weight: 700; color: var(--ink); }

/* history groups */
.grp {
  display: flex; align-items: center; gap: 8px; font-weight: 800; font-size: 0.74rem;
  letter-spacing: .07em; text-transform: uppercase; color: var(--muted);
  margin: 18px 0 8px 0; animation: fadeIn .4s ease both;
}
.grp-dot { width: 9px; height: 9px; border-radius: 50%; }
.grp-count {
  margin-left: auto; background: #fff; border: 1px solid var(--line); border-radius: 999px;
  padding: 0 8px; font-size: 0.7rem; color: var(--ink);
}
.hist-sub { color: var(--muted); font-size: 0.72rem; margin: -6px 0 10px 4px; line-height: 1.3; animation: fadeIn .5s ease both; }
.hist-empty {
  border: 1.5px dashed var(--line); border-radius: 12px; padding: 14px; color: var(--muted);
  font-size: 0.84rem; text-align: center; margin-top: 8px; animation: fadeIn .5s ease both;
}

/* history buttons (keys: chat_<id>, chatactive_<id>, del_<id>) */
[class*="st-key-chat_"] button, [class*="st-key-chatactive_"] button {
  justify-content: flex-start !important; text-align: left !important;
  background: #fff !important; color: var(--ink) !important; border: 1px solid var(--line) !important;
  border-radius: 10px !important; font-weight: 700 !important; font-size: 0.82rem !important;
  padding: 0.45rem 0.7rem !important; min-height: 0 !important; box-shadow: none !important;
  animation: slideLeft .45s var(--ease) both;
}
[class*="st-key-chat_"] button p, [class*="st-key-chatactive_"] button p {
  white-space: nowrap; overflow: hidden; text-overflow: ellipsis; font-size: 0.82rem;
}
[class*="st-key-chat_"] button:hover {
  border-color: var(--blue) !important; background: var(--soft) !important; color: var(--ink) !important;
  transform: translateX(4px); box-shadow: 0 6px 14px rgba(33,80,200,0.12) !important;
}
[class*="st-key-chatactive_"] button {
  border-left: 4px solid var(--red) !important; background: #fff !important;
  box-shadow: 0 6px 16px rgba(33,80,200,0.14) !important;
}
[class*="st-key-del_"] button {
  background: transparent !important; border: 0 !important; color: var(--muted) !important;
  min-height: 0 !important; padding: 0.35rem 0 !important; box-shadow: none !important;
}
[class*="st-key-del_"] button:hover { color: var(--red) !important; background: #FDECEE !important; transform: none; }

/* suggestion chips */
[class*="st-key-sug_"] button {
  background: #fff !important; color: var(--ink) !important; border: 1.5px solid var(--line) !important;
  border-bottom: 3px solid var(--yellow) !important; border-radius: 14px !important;
  font-weight: 700 !important; font-size: 0.88rem !important; padding: 0.8rem 1rem !important;
  text-align: left !important; justify-content: flex-start !important; height: 100%;
  animation: fadeUp .55s var(--ease) both;
}
[class*="st-key-sug_"] button:hover {
  border-color: var(--blue) !important; background: var(--soft) !important; color: var(--blue) !important;
  transform: translateY(-3px); box-shadow: 0 10px 22px rgba(33,80,200,0.15) !important;
}

/* ---------- File uploader ---------- */
[data-testid="stFileUploader"] section {
  border: 2px dashed var(--blue); background: var(--soft); border-radius: 14px; padding: 1.4rem;
  transition: border-color .25s, transform .25s var(--ease), box-shadow .25s;
}
[data-testid="stFileUploader"] section:hover { border-color: var(--red); transform: translateY(-2px); box-shadow: 0 10px 24px rgba(33,80,200,0.12); }

/* ---------- Chat ---------- */
[data-testid="stChatMessage"] {
  border: 1px solid var(--line); border-radius: 12px; padding: 1rem 1.1rem;
  background: #fff; margin-bottom: 0.8rem; animation: fadeUp .45s var(--ease) both;
  transition: box-shadow .25s, transform .25s var(--ease);
}
[data-testid="stChatMessage"]:hover { box-shadow: 0 8px 22px rgba(15,27,61,0.08); transform: translateY(-1px); }
[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]) {
  background: var(--soft); border-left: 4px solid var(--blue); animation-name: slideRight;
}
[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarAssistant"]) {
  border-left: 4px solid var(--red); animation-name: slideLeft;
}
[data-testid="stChatInput"] { border: 1.5px solid var(--line); border-radius: 12px; transition: border-color .2s, box-shadow .2s; }
[data-testid="stChatInput"]:focus-within { border-color: var(--blue); box-shadow: 0 0 0 3px rgba(33,80,200,0.15); }

[data-testid="stExpander"] { border: 1px solid var(--line); border-radius: 12px; background: #fff; animation: fadeIn .5s ease both; }
[data-testid="stExpander"] summary { font-weight: 700; color: var(--blue); }

.page-chip {
  display: inline-block; background: var(--yellow); color: var(--ink); font-weight: 800;
  font-size: 0.78rem; padding: 2px 10px; border-radius: 999px; margin-bottom: 6px;
  animation: popIn .4s var(--ease) both;
}
.msg-time { color: var(--muted); font-size: 0.72rem; margin-top: 4px; }

/* typing indicator */
.typing { display: inline-flex; align-items: center; gap: 6px; padding: 6px 2px; }
.typing span.d { width: 9px; height: 9px; border-radius: 50%; background: var(--blue); animation: bounce 1.2s infinite; }
.typing span.d:nth-child(2) { background: var(--red); animation-delay: .15s; }
.typing span.d:nth-child(3) { background: var(--yellow); animation-delay: .3s; }
.typing .lbl { color: var(--muted); font-size: 0.85rem; margin-left: 8px; font-weight: 600; }

/* ---------- App header ---------- */
.app-header {
  position: relative; display: flex; align-items: center; justify-content: space-between;
  background: linear-gradient(120deg, var(--blue-dark), var(--blue), #3A6BE0, var(--blue-dark));
  background-size: 300% 300%; animation: gradientShift 12s ease infinite, fadeUp .6s var(--ease) both;
  color: #fff; border-radius: 16px; padding: 20px 26px; margin-bottom: 22px; overflow: hidden;
}
.app-header::after {
  content: ""; position: absolute; left: 0; right: 0; bottom: 0; height: 5px; transform-origin: left;
  background: linear-gradient(90deg, var(--red) 0 70%, var(--yellow) 70% 100%);
  animation: barGrow 1s var(--ease) .3s both;
}
.app-header .title { font-weight: 800; font-size: 1.35rem; line-height: 1.2; }
.app-header .sub { opacity: 0.85; font-size: 0.9rem; word-break: break-word; }
.chip {
  display: inline-flex; align-items: center; gap: 8px; background: rgba(255,255,255,0.14);
  border: 1px solid rgba(255,255,255,0.35); border-radius: 999px; padding: 6px 14px;
  font-weight: 700; font-size: 0.85rem; white-space: nowrap;
}
.chip .dot { width: 9px; height: 9px; border-radius: 50%; background: var(--yellow); animation: pulseDot 2s infinite; }

/* ---------- Stat cards ---------- */
.stats { display: flex; gap: 14px; margin: 6px 0 18px 0; flex-wrap: wrap; }
.stat {
  flex: 1 1 160px; background: #fff; border: 1px solid var(--line); border-top: 4px solid var(--blue);
  border-radius: 12px; padding: 12px 16px; animation: fadeUp .55s var(--ease) both;
  transition: transform .25s var(--ease), box-shadow .25s;
}
.stat:hover { transform: translateY(-4px); box-shadow: 0 12px 26px rgba(33,80,200,0.14); }
.stat:nth-child(1) { animation-delay: .05s; }
.stat:nth-child(2) { border-top-color: var(--red); animation-delay: .15s; }
.stat:nth-child(3) { border-top-color: var(--yellow); animation-delay: .25s; }
.stat:nth-child(4) { border-top-color: var(--blue-dark); animation-delay: .35s; }
.stat .s-label { color: var(--muted); font-size: 0.78rem; font-weight: 600; }
.stat .s-value { font-weight: 800; font-size: 1.15rem; word-break: break-word; }

/* ---------- Welcome / hero ---------- */
.orb { position: absolute; border-radius: 50%; filter: blur(2px); opacity: 0.55; pointer-events: none; }
.orb.o1 { width: 220px; height: 220px; right: -70px; top: -60px; background: radial-gradient(circle at 30% 30%, rgba(255,199,44,0.65), transparent 70%); animation: drift 9s ease-in-out infinite; }
.orb.o2 { width: 180px; height: 180px; left: -60px; bottom: 60px; background: radial-gradient(circle at 30% 30%, rgba(214,40,57,0.6), transparent 70%); animation: drift 11s ease-in-out -3s infinite; }
.orb.o3 { width: 110px; height: 110px; right: 40px; bottom: 120px; background: radial-gradient(circle at 30% 30%, rgba(255,255,255,0.4), transparent 70%); animation: drift 7s ease-in-out -1s infinite; }

.hero {
  position: relative; overflow: hidden; border-radius: 20px; padding: 44px 40px 52px 40px; color: #fff; margin-bottom: 22px;
  background: linear-gradient(150deg, var(--blue-dark), var(--blue), #3A6BE0, var(--blue-dark));
  background-size: 300% 300%; animation: gradientShift 14s ease infinite, fadeUp .7s var(--ease) both;
}
.hero::after {
  content: ""; position: absolute; left: 0; right: 0; bottom: 0; height: 7px; transform-origin: left;
  background: linear-gradient(90deg, var(--red) 0 72%, var(--yellow) 72% 100%);
  animation: barGrow 1.2s var(--ease) .5s both;
}
.hero > *:not(.orb) { position: relative; z-index: 1; }
.hero h1 { color: #fff; font-size: 2.2rem; line-height: 1.2; font-weight: 800; margin: 0 0 12px 0; padding: 0; animation: fadeUp .8s var(--ease) .15s both; }
.hero p { color: rgba(255,255,255,0.9); font-size: 1.02rem; line-height: 1.6; max-width: 620px; margin: 0; animation: fadeUp .8s var(--ease) .3s both; }
.hero .badge {
  display: inline-block; background: var(--yellow); color: var(--ink); font-weight: 800; font-size: 0.78rem;
  padding: 4px 12px; border-radius: 999px; margin-bottom: 14px; animation: popIn .6s var(--ease) both;
}

.type-grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(190px, 1fr)); gap: 14px; margin: 0 0 22px 0; }
.type-card {
  background: #fff; border: 1px solid var(--line); border-radius: 14px; padding: 16px 18px;
  border-top: 4px solid var(--blue); animation: fadeUp .6s var(--ease) both;
  transition: transform .25s var(--ease), box-shadow .25s;
}
.type-card:nth-child(1) { border-top-color: var(--red);  animation-delay: .1s; }
.type-card:nth-child(2) { border-top-color: var(--blue); animation-delay: .2s; }
.type-card:nth-child(3) { border-top-color: var(--yellow); animation-delay: .3s; }
.type-card:nth-child(4) { border-top-color: var(--blue-dark); animation-delay: .4s; }
.type-card:hover { transform: translateY(-5px); box-shadow: 0 14px 28px rgba(33,80,200,0.15); }
.type-card:hover .ic { animation: wiggle .6s ease; }
.type-card .ic { font-size: 1.8rem; display: inline-block; }
.type-card b { display: block; font-size: 1rem; margin-top: 6px; }
.type-card span { color: var(--muted); font-size: 0.82rem; }

.sect-label { font-weight: 800; color: var(--blue); font-size: 0.95rem; margin: 6px 0 10px 0; }

.empty {
  border: 1px solid var(--line); border-left: 5px solid var(--yellow); background: var(--soft);
  border-radius: 12px; padding: 18px 20px; margin-top: 8px; animation: fadeUp .6s var(--ease) .15s both;
}
.empty b { color: var(--blue); }

[data-testid="stAlert"] { border-radius: 12px; animation: popIn .4s var(--ease) both; }

@media (prefers-reduced-motion: reduce) {
  *, *::before, *::after { animation-duration: 0.001ms !important; animation-iteration-count: 1 !important; transition-duration: 0.001ms !important; }
}
</style>
"""

LOGIN_CSS = """
<style>
[data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"],
[data-testid="collapsedControl"] { display: none !important; }

.login-panel {
  position: relative; min-height: 600px; border-radius: 18px; padding: 40px 38px 50px 38px; color: #fff;
  background: linear-gradient(150deg, var(--blue-dark), var(--blue), #3A6BE0, var(--blue-dark));
  background-size: 300% 300%; overflow: hidden;
  animation: gradientShift 14s ease infinite, slideLeft .8s var(--ease) both;
}
.login-panel::after {
  content: ""; position: absolute; left: 0; right: 0; bottom: 0; height: 8px; transform-origin: left;
  background: linear-gradient(90deg, var(--red) 0 72%, var(--yellow) 72% 100%);
  animation: barGrow 1.2s var(--ease) .6s both;
}
.login-panel > *:not(.orb) { position: relative; z-index: 1; }
.login-panel .brand { animation: fadeUp .7s var(--ease) .1s both; }
.login-panel .brand-name { color: #fff; }
.login-panel h1 {
  color: #fff; font-size: 2.1rem; line-height: 1.2; font-weight: 800; margin: 44px 0 14px 0; padding: 0;
  animation: fadeUp .8s var(--ease) .25s both;
}
.login-panel .lead { color: rgba(255,255,255,0.88); font-size: 1.02rem; line-height: 1.6; max-width: 460px; animation: fadeUp .8s var(--ease) .4s both; }
.login-panel ul { list-style: none; padding: 0; margin: 28px 0 0 0; }
.login-panel li {
  position: relative; padding: 0 0 18px 30px; color: rgba(255,255,255,0.9); font-size: 0.96rem; line-height: 1.5;
  animation: fadeUp .7s var(--ease) both;
}
.login-panel li:nth-child(1) { animation-delay: .6s; }
.login-panel li:nth-child(2) { animation-delay: .78s; }
.login-panel li:nth-child(3) { animation-delay: .96s; }
.login-panel li::before {
  content: ""; position: absolute; left: 0; top: 5px; width: 14px; height: 14px;
  background: var(--yellow); border-radius: 4px; animation: pulseDot 2.4s infinite;
}
.login-panel li b { display: block; color: #fff; font-size: 1rem; }

.login-spacer { height: 36px; }
.login-title { font-weight: 800; font-size: 1.6rem; margin: 0 0 4px 0; color: var(--ink); animation: fadeUp .6s var(--ease) .3s both; }
.login-sub { color: var(--muted); font-size: 0.9rem; margin-bottom: 14px; animation: fadeUp .6s var(--ease) .4s both; }

[data-testid="stForm"] {
  background: #fff; border: 1px solid var(--line); border-top: 5px solid var(--red);
  border-radius: 16px; padding: 28px 28px 22px 28px; box-shadow: 0 10px 30px rgba(33,80,200,0.10);
  animation: slideRight .8s var(--ease) .15s both;
  transition: box-shadow .3s var(--ease), transform .3s var(--ease);
}
[data-testid="stForm"]:hover { box-shadow: 0 18px 44px rgba(33,80,200,0.18); transform: translateY(-3px); }
[data-testid="stForm"] [data-testid="stTextInput"] { animation: fadeUp .6s var(--ease) .5s both; }
[data-testid="stForm"] [data-testid="stFormSubmitButton"] { animation: fadeUp .6s var(--ease) .65s both; }
[data-testid="stAlert"][data-baseweb="notification"], .stAlert { animation: shake .5s ease both; }

@media (max-width: 800px) { .login-panel { min-height: 0; padding: 28px 24px 40px 24px; } }
</style>
"""

LOGO_SVG = (
    '<svg width="40" height="40" viewBox="0 0 40 40" xmlns="http://www.w3.org/2000/svg" aria-hidden="true">'
    '<rect width="40" height="40" rx="10" fill="#FFFFFF"/>'
    '<path d="M12 8h11l6 6v18H12z" fill="#2150C8"/>'
    '<path d="M23 8v6h6z" fill="#FFC72C"/>'
    '<rect x="15.5" y="19" width="10" height="2" rx="1" fill="#FFFFFF"/>'
    '<rect x="15.5" y="24" width="7" height="2" rx="1" fill="#D62839"/>'
    "</svg>"
)
LOGO_SVG_SIDEBAR = LOGO_SVG.replace('<rect width="40" height="40" rx="10" fill="#FFFFFF"/>',
                                    '<rect width="40" height="40" rx="10" fill="#FFFFFF" stroke="#DCE3F5"/>')

LOGIN_PANEL_HTML = (
    '<div class="login-panel">'
    '<span class="orb o1"></span><span class="orb o2"></span><span class="orb o3"></span>'
    f'<div class="brand">{LOGO_SVG}<span class="brand-name">Document RAG Agent</span></div>'
    "<h1>Ask your documents. Get answers with clear citations.</h1>"
    '<p class="lead">Upload a PDF, Word, CSV or Excel file and ask questions in plain language. '
    "Every answer is based only on your file and shows where it came from.</p>"
    "<ul>"
    "<li><b>Grounded answers</b>Responses use only the content of your file.</li>"
    "<li><b>Clear citations</b>See the page, section or rows behind each answer.</li>"
    "<li><b>Your key stays private</b>It is kept in this browser session only and is never saved.</li>"
    "</ul>"
    "</div>"
)

THINKING_HTML = (
    '<div class="typing"><span class="d"></span><span class="d"></span><span class="d"></span>'
    '<span class="lbl">Searching your document…</span></div>'
)

st.markdown(BASE_CSS, unsafe_allow_html=True)

# =============================================================================
# File types (used for grouping history in the sidebar)
# =============================================================================
SUPPORTED_TYPES = ["pdf", "docx", "doc", "csv", "xlsx", "xls"]

# key -> (label, icon, colour)
TYPE_META = {
    "pdf":   ("PDF",   "📕", "#D62839"),
    "word":  ("Word",  "📘", "#2150C8"),
    "csv":   ("CSV",   "📒", "#FFC72C"),
    "excel": ("Excel", "📊", "#16399A"),
}
TYPE_ORDER = ["pdf", "word", "csv", "excel"]


def ftype_of(name: str) -> str:
    ext = name.rsplit(".", 1)[-1].lower()
    return {"pdf": "pdf", "docx": "word", "doc": "word", "csv": "csv",
            "xlsx": "excel", "xls": "excel"}.get(ext, "pdf")


# =============================================================================
# Chat history storage (JSON file + saved copies of uploaded files)
#
# chat_history.json:
# {
#   "version": 2,
#   "users": {
#     "<user_id>": {                       # sha256 of the API key (key itself is never stored)
#       "chats": {
#         "<chat_id>": {
#           "id", "file", "type", "path", "title",
#           "created", "updated",
#           "messages": [ {role, content, time, sources?}, ... ]
#         }
#       }
#     }
#   }
# }
# The uploaded file is saved in ./saved_files/<user_id>/ so an old chat can be
# re-opened from the sidebar and you can keep asking questions.
# =============================================================================
HISTORY_FILE = Path(os.environ.get("CHAT_HISTORY_FILE", Path(__file__).with_name("chat_history.json")))
UPLOAD_DIR = Path(os.environ.get("SAVED_FILES_DIR", Path(__file__).with_name("saved_files")))
_HISTORY_LOCK = threading.Lock()


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _read_all() -> dict:
    try:
        with open(HISTORY_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict) and isinstance(data.get("users"), dict):
            return data
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        pass
    return {"version": 2, "users": {}}


def _write_all(data: dict) -> None:
    """Atomic write so a crash never leaves a half-written JSON file."""
    data["version"] = 2
    tmp = HISTORY_FILE.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    os.replace(tmp, HISTORY_FILE)


def user_id(api_key: str) -> str:
    return hashlib.sha256(api_key.encode("utf-8")).hexdigest()[:16]


def _chats(data: dict, uid: str) -> dict:
    return data["users"].setdefault(uid, {}).setdefault("chats", {})


def create_chat(uid: str, filename: str, file_bytes: bytes) -> str:
    """Save a copy of the file and register a brand-new chat for it."""
    ext = filename.rsplit(".", 1)[-1].lower()
    path = UPLOAD_DIR / uid / f"{hashlib.sha1(file_bytes).hexdigest()[:16]}.{ext}"
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        path.write_bytes(file_bytes)
    cid = uuid.uuid4().hex[:12]
    with _HISTORY_LOCK:
        data = _read_all()
        _chats(data, uid)[cid] = {
            "id": cid, "file": filename, "type": ftype_of(filename), "path": str(path),
            "title": "", "created": _now(), "updated": _now(), "messages": [],
        }
        _write_all(data)
    return cid


def get_chat(uid: str, cid: str):
    with _HISTORY_LOCK:
        data = _read_all()
    return data["users"].get(uid, {}).get("chats", {}).get(cid)


def save_messages(uid: str, cid: str, messages: list) -> None:
    try:
        with _HISTORY_LOCK:
            data = _read_all()
            entry = _chats(data, uid).get(cid)
            if entry is None:
                return
            entry["messages"] = messages
            entry["updated"] = _now()
            if not entry.get("title"):
                first = next((m["content"] for m in messages if m["role"] == "user"), "")
                entry["title"] = first.strip().replace("\n", " ")[:60]
            _write_all(data)
    except OSError as e:
        st.warning(f"Chat history could not be saved: {e}")


def delete_chat(uid: str, cid: str) -> None:
    try:
        with _HISTORY_LOCK:
            data = _read_all()
            chats = _chats(data, uid)
            entry = chats.pop(cid, None)
            _write_all(data)
            # remove the saved file only if no other chat uses it
            if entry and not any(c.get("path") == entry.get("path") for c in chats.values()):
                try:
                    Path(entry["path"]).unlink(missing_ok=True)
                except OSError:
                    pass
    except OSError as e:
        st.warning(f"Chat could not be deleted: {e}")


def list_chats(uid: str) -> list:
    with _HISTORY_LOCK:
        data = _read_all()
    chats = data["users"].get(uid, {}).get("chats", {})
    items = [
        {"id": c.get("id", k), "file": c.get("file", "file"), "type": c.get("type", "pdf"),
         "title": c.get("title", ""), "count": len(c.get("messages", [])),
         "updated": c.get("updated", "")}
        for k, c in chats.items()
    ]
    return sorted(items, key=lambda x: x["updated"], reverse=True)


def export_history(uid: str) -> str:
    with _HISTORY_LOCK:
        data = _read_all()
    return json.dumps(data["users"].get(uid, {"chats": {}}), ensure_ascii=False, indent=2)


def export_chat_md(file_name: str, messages: list) -> str:
    out = [f"# Chat about {file_name}", ""]
    for m in messages:
        who = "You" if m["role"] == "user" else "Assistant"
        out.append(f"**{who}** ({m.get('time', '').replace('T', ' ')})")
        out.append("")
        out.append(m["content"])
        out.append("")
    return "\n".join(out)


# =============================================================================
# Session helpers / callbacks
# =============================================================================
CHAT_STATE_KEYS = ("active_chat", "vectorstore", "index_sig", "messages", "messages_for",
                   "doc_info", "pending_q")


def reset_chat_state() -> None:
    for k in CHAT_STATE_KEYS:
        st.session_state.pop(k, None)


def new_chat_cb() -> None:
    reset_chat_state()


def open_chat_cb(cid: str) -> None:
    reset_chat_state()
    st.session_state.active_chat = cid


def delete_chat_cb(cid: str) -> None:
    delete_chat(user_id(st.session_state.api_key), cid)
    if st.session_state.get("active_chat") == cid:
        reset_chat_state()


def ask_cb(text: str) -> None:
    st.session_state.pending_q = text


def short(text: str, n: int) -> str:
    return text if len(text) <= n else text[: n - 1] + "…"


# =============================================================================
# API key verification
# =============================================================================
def verify_api_key(key: str) -> tuple[bool, str]:
    """Make a lightweight authenticated call to confirm the key works."""
    key = key.strip()
    if not key:
        return False, "Enter your API key to continue."
    if not key.startswith("sk-"):
        return False, "This doesn't look like an OpenAI API key. Keys start with sk-."
    try:
        OpenAI(api_key=key, timeout=15, max_retries=1).models.list()
        return True, ""
    except AuthenticationError:
        return False, "The key was rejected. Check that you copied it completely and that it hasn't been revoked."
    except PermissionDeniedError:
        return False, "This key doesn't have permission to access the API. Check its project and permissions."
    except RateLimitError:
        return False, "The key is valid but has hit a rate or quota limit. Check your billing and usage, then try again."
    except APIConnectionError:
        return False, "Couldn't reach OpenAI. Check your internet connection and try again."
    except APIStatusError as e:
        return False, f"OpenAI returned an error ({e.status_code}). Please try again in a moment."
    except Exception as e:  # noqa: BLE001
        return False, f"Unexpected error while verifying the key: {e}"


def mask_key(key: str) -> str:
    return f"{key[:3]}…{key[-4:]}" if len(key) > 8 else "••••"


def sign_out() -> None:
    # Saved chat history stays on disk; only the in-memory session is cleared.
    reset_chat_state()
    for k in ("authenticated", "api_key"):
        st.session_state.pop(k, None)
    st.rerun()


# =============================================================================
# Login page
# =============================================================================
def render_login() -> None:
    st.markdown(LOGIN_CSS, unsafe_allow_html=True)
    left, right = st.columns([1.15, 1], gap="large")

    with left:
        st.markdown(LOGIN_PANEL_HTML, unsafe_allow_html=True)

    with right:
        st.markdown('<div class="login-spacer"></div>', unsafe_allow_html=True)
        with st.form("api_key_form", clear_on_submit=False):
            st.markdown('<div class="login-title">Enter your API key</div>', unsafe_allow_html=True)
            st.markdown('<div class="login-sub">Verified securely. Never stored on disk.</div>',
                        unsafe_allow_html=True)
            key = st.text_input("OpenAI API key", type="password", placeholder="sk-...", label_visibility="collapsed")
            submitted = st.form_submit_button(
                "Verify and continue", type="primary", use_container_width=True
            )

        if submitted:
            with st.spinner("Verifying your key…"):
                ok, message = verify_api_key(key)
            if ok:
                st.session_state.api_key = key.strip()
                st.session_state.authenticated = True
                st.success("Key verified. Loading your workspace…")
                time.sleep(0.8)  # let the success animation play
                st.rerun()
            else:
                st.error(message)


# =============================================================================
# File loaders: PDF, Word, CSV, Excel
# =============================================================================
MAX_ROWS = 20000
ROWS_PER_BLOCK = 25


def _load_pdf(data: bytes):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        docs = PyPDFLoader(tmp_path).load()
    finally:
        os.unlink(tmp_path)
    for d in docs:
        p = d.metadata.get("page")
        d.metadata["label"] = f"Page {p + 1}" if isinstance(p, int) else "Page ?"
    return docs, "Pages", len(docs), ""


def _load_docx(data: bytes):
    if not data.startswith(b"PK"):
        raise ValueError(
            "Old .doc files can't be read directly. Open it in Word, choose "
            "Save As → .docx, and upload the new file."
        )
    try:
        from docx import Document as DocxDocument
    except ImportError:
        raise ImportError("python-docx package not found. Install it with: python -m pip install python-docx")

    doc = DocxDocument(io.BytesIO(data))
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    sections, buf, size = [], [], 0
    for line in lines:
        buf.append(line)
        size += len(line)
        if size >= 3000:
            sections.append("\n".join(buf))
            buf, size = [], 0
    if buf:
        sections.append("\n".join(buf))

    docs = [
        Document(page_content=t, metadata={"label": f"Section {i}"})
        for i, t in enumerate(sections, 1)
    ]
    return docs, "Sections", len(docs), ""


def _table_to_docs(df, prefix: str = ""):
    """Turn a table into row-block documents that keep the column names."""
    df = df.fillna("").astype(str)
    df = df[(df != "").any(axis=1)].reset_index(drop=True)
    cols = [str(c) for c in df.columns]
    docs = []
    for start in range(0, len(df), ROWS_PER_BLOCK):
        block = df.iloc[start:start + ROWS_PER_BLOCK]
        lines = []
        for i, values in enumerate(block.values.tolist()):
            parts = [f"{c}: {v}" for c, v in zip(cols, values) if str(v).strip()]
            lines.append(f"Row {start + i + 1}: " + " | ".join(parts))
        label = f"{prefix}Rows {start + 1}-{start + len(block)}"
        docs.append(Document(page_content="\n".join(lines), metadata={"label": label}))
    return docs, len(df)


def _load_csv(data: bytes):
    import pandas as pd

    df = None
    for enc in ("utf-8-sig", "latin-1"):
        try:
            df = pd.read_csv(io.BytesIO(data), dtype=str, keep_default_na=False,
                             sep=None, engine="python", encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    if df is None:
        raise ValueError("Couldn't decode this CSV file. Save it as UTF-8 and try again.")
    note = ""
    if len(df) > MAX_ROWS:
        df = df.head(MAX_ROWS)
        note = f"Large file: only the first {MAX_ROWS:,} rows were indexed."
    docs, n_rows = _table_to_docs(df)
    return docs, "Rows", n_rows, note


def _load_excel(data: bytes):
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str, keep_default_na=False)
    docs, total, note = [], 0, ""
    for name, df in sheets.items():
        remaining = MAX_ROWS - total
        if remaining <= 0:
            note = f"Large file: only the first {MAX_ROWS:,} rows were indexed."
            break
        if len(df) > remaining:
            df = df.head(remaining)
            note = f"Large file: only the first {MAX_ROWS:,} rows were indexed."
        sheet_docs, n_rows = _table_to_docs(df, prefix=f"{name}, ")
        docs.extend(sheet_docs)
        total += n_rows
    return docs, "Rows", total, note


def load_documents(data: bytes, filename: str):
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return _load_pdf(data)
    if ext in ("docx", "doc"):
        return _load_docx(data)
    if ext == "csv":
        return _load_csv(data)
    if ext in ("xlsx", "xls"):
        return _load_excel(data)
    raise ValueError(f"Unsupported file type: .{ext}")


# =============================================================================
# RAG pieces
# =============================================================================
RAG_PROMPT = ChatPromptTemplate.from_messages([
    ("system",
     "You are a helpful assistant that answers questions strictly using the "
     "provided context from an uploaded document (PDF, Word, CSV or Excel).\n"
     "Rules:\n"
     "1. Answer ONLY from the context below.\n"
     "2. Treat \"you\" in the question as referring to the document. For broad or "
     "general questions (e.g. \"what do you know about X\", \"what is this about\", "
     "\"summarize\"), give a clear summary of everything the context says on that "
     "topic, even if it only covers part of it.\n"
     "3. Only if the context has nothing relevant at all, say "
     "\"I couldn't find that in the document.\"\n"
     "4. Cite where you found it using the labels in the context, "
     "e.g. (Page 3), (Section 2) or (Sheet1, Rows 1-25).\n"
     "5. The conversation history is only for understanding follow-up questions; "
     "never use it as a source of facts.\n\n"
     "Conversation history:\n{history}\n\n"
     "Context:\n{context}"),
    ("human", "{question}"),
])


def format_docs(docs) -> str:
    return "\n\n".join(f"[{d.metadata.get('label', 'Source')}]\n{d.page_content}" for d in docs)


def format_history(messages: list, turns: int = 3) -> str:
    recent = messages[-turns * 2:]
    if not recent:
        return "(none)"
    return "\n".join(f"{m['role'].capitalize()}: {m['content'][:600]}" for m in recent)


def build_vectorstore(data: bytes, filename: str, chunk_size: int, overlap: int, api_key: str):
    """Load the file, split it into chunks, embed, and store in FAISS."""
    docs, unit_name, unit_count, note = load_documents(data, filename)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = [c for c in splitter.split_documents(docs) if c.page_content.strip()]
    if not chunks:
        raise ValueError(
            "No readable content found. If this is a scanned PDF, try a text-based version."
        )
    embeddings = OpenAIEmbeddings(model="text-embedding-3-small", api_key=api_key)
    vs = FAISS.from_documents(chunks, embeddings)
    return vs, unit_name, unit_count, len(chunks), note


def retrieve(question: str, vectorstore, k: int):
    return vectorstore.as_retriever(search_kwargs={"k": k}).invoke(question)


def stream_answer(question: str, docs, history: str, model: str, api_key: str):
    """Yield the answer token by token so the UI can stream it."""
    llm = ChatOpenAI(model=model, temperature=0, api_key=api_key)
    chain = RAG_PROMPT | llm | StrOutputParser()
    yield from chain.stream(
        {"context": format_docs(docs), "question": question, "history": history}
    )


def docs_to_sources(docs) -> list:
    return [{"label": d.metadata.get("label", "Source"), "text": d.page_content[:500]} for d in docs]


def render_sources(sources) -> None:
    with st.expander(f"Sources ({len(sources)} retrieved passages)"):
        for s in sources:
            st.markdown(
                f'<span class="page-chip">{html.escape(str(s["label"]))}</span>',
                unsafe_allow_html=True,
            )
            st.text(s["text"])
            st.divider()


def render_message(msg: dict) -> None:
    with st.chat_message(msg["role"], avatar="👤" if msg["role"] == "user" else "📄"):
        st.markdown(msg["content"])
        if msg.get("sources"):
            render_sources(msg["sources"])
        if msg.get("time"):
            st.markdown(f'<div class="msg-time">{html.escape(msg["time"].replace("T", " "))}</div>',
                        unsafe_allow_html=True)


# =============================================================================
# Sidebar
# =============================================================================
def render_sidebar(uid: str, api_key: str):
    """Draws the sidebar and returns (model_name, chunk_size, chunk_overlap, top_k)."""
    active_id = st.session_state.get("active_chat")

    with st.sidebar:
        st.markdown(
            f'<div class="side-brand">{LOGO_SVG_SIDEBAR}<b>Document RAG Agent</b></div>',
            unsafe_allow_html=True,
        )
        st.button("＋  New chat", key="new_chat_btn", type="primary",
                  use_container_width=True, on_click=new_chat_cb)

        st.markdown('<div class="side-title" style="margin-top:1.1rem">History</div>',
                    unsafe_allow_html=True)
        search = st.text_input("Search chats", placeholder="🔍 Search chats or files…",
                               label_visibility="collapsed", key="chat_search")

        chats = list_chats(uid)
        if search.strip():
            s = search.strip().lower()
            chats = [c for c in chats if s in c["file"].lower() or s in c["title"].lower()]

        if not chats:
            msg = "No chats match your search." if search.strip() else (
                "No chats yet.<br>Upload a file to start your first one.")
            st.markdown(f'<div class="hist-empty">{msg}</div>', unsafe_allow_html=True)
        else:
            groups = {}
            for c in chats:
                groups.setdefault(c["type"], []).append(c)
            for t in TYPE_ORDER:
                items = groups.get(t)
                if not items:
                    continue
                label, icon, color = TYPE_META[t]
                st.markdown(
                    f'<div class="grp"><span class="grp-dot" style="background:{color}"></span>'
                    f'{icon} {label}<span class="grp-count">{len(items)}</span></div>',
                    unsafe_allow_html=True,
                )
                for c in items:
                    is_active = c["id"] == active_id
                    col_a, col_b = st.columns([6, 1], gap="small", vertical_alignment="center")
                    with col_a:
                        st.button(
                            short(c["file"], 28),
                            key=f"{'chatactive' if is_active else 'chat'}_{c['id']}",
                            on_click=open_chat_cb, args=(c["id"],),
                            use_container_width=True,
                            help=f'{c["file"]}  ·  {c["updated"].replace("T", " ")}',
                        )
                    with col_b:
                        st.button("✕", key=f"del_{c['id']}", on_click=delete_chat_cb,
                                  args=(c["id"],), help="Delete this chat")
                    sub = short(c["title"], 38) if c["title"] else "New chat"
                    st.markdown(
                        f'<div class="hist-sub">{html.escape(sub)} · {c["count"]} msgs</div>',
                        unsafe_allow_html=True,
                    )

        st.markdown("<div style='height:14px'></div>", unsafe_allow_html=True)
        with st.expander("⚙️  Settings"):
            model_name = st.selectbox("Chat model", ["gpt-4o-mini", "gpt-4o", "gpt-4.1-mini"])
            chunk_size = st.slider("Chunk size", 500, 2000, 1000, step=100)
            chunk_overlap = st.slider("Chunk overlap", 0, 400, 150, step=50)
            top_k = st.slider("Retrieved chunks (k)", 2, 10, 4)
            st.caption("Changing chunk size or overlap re-indexes the file.")

        with st.expander("👤  Account"):
            st.markdown(
                '<div class="key-box"><div class="k-label">API key (verified)</div>'
                f'<div class="k-value">{mask_key(api_key)}</div></div>',
                unsafe_allow_html=True,
            )
            st.download_button(
                "Download all history (JSON)",
                data=export_history(uid),
                file_name="chat_history.json",
                mime="application/json",
                use_container_width=True,
            )
            if st.button("Change API key", use_container_width=True):
                sign_out()

    return model_name, chunk_size, chunk_overlap, top_k


# =============================================================================
# Main views
# =============================================================================
def render_welcome(uid: str) -> None:
    st.markdown(
        '<div class="hero">'
        '<span class="orb o1"></span><span class="orb o2"></span><span class="orb o3"></span>'
        '<span class="badge">✨ New chat</span>'
        "<h1>What would you like to explore today?</h1>"
        "<p>Upload a document and start asking questions. Every answer is grounded in your file "
        "and comes with the exact page, section or rows it was taken from.</p>"
        "</div>"
        '<div class="type-grid">'
        '<div class="type-card"><span class="ic">📕</span><b>PDF</b><span>Reports, papers, books</span></div>'
        '<div class="type-card"><span class="ic">📘</span><b>Word</b><span>.docx documents</span></div>'
        '<div class="type-card"><span class="ic">📒</span><b>CSV</b><span>Tables and exports</span></div>'
        '<div class="type-card"><span class="ic">📊</span><b>Excel</b><span>.xlsx / .xls, all sheets</span></div>'
        "</div>",
        unsafe_allow_html=True,
    )
    st.markdown('<div class="sect-label">Upload your file</div>', unsafe_allow_html=True)
    uploaded = st.file_uploader(
        "Upload a file (PDF, Word, CSV or Excel)", type=SUPPORTED_TYPES, label_visibility="collapsed"
    )
    if uploaded is not None:
        with st.spinner("Saving your file…"):
            cid = create_chat(uid, uploaded.name, uploaded.getvalue())
        st.session_state.active_chat = cid
        st.rerun()

    st.markdown(
        '<div class="empty"><b>Your chats are saved automatically.</b><br>'
        "They appear on the left, grouped by file type. Click any of them later to continue "
        "where you stopped, with no need to upload the file again.</div>",
        unsafe_allow_html=True,
    )


def ensure_index(chat: dict, cid: str, api_key: str, chunk_size: int, chunk_overlap: int) -> bool:
    """Build (or rebuild) the vector index for the open chat. Returns True when ready."""
    sig = (cid, chunk_size, chunk_overlap)
    if st.session_state.get("index_sig") == sig and "vectorstore" in st.session_state:
        return True
    try:
        data = Path(chat["path"]).read_bytes()
    except OSError:
        st.error("The saved copy of this file is missing. Start a new chat and upload it again.")
        return False
    try:
        with st.spinner("Reading and indexing your file…"):
            vs, unit_name, unit_count, n_chunks, note = build_vectorstore(
                data, chat["file"], chunk_size, chunk_overlap, api_key
            )
        st.session_state.vectorstore = vs
        st.session_state.doc_info = {
            "name": chat["file"], "type": chat["type"], "unit_name": unit_name,
            "units": unit_count, "chunks": n_chunks, "note": note,
        }
        st.session_state.index_sig = sig
        return True
    except AuthenticationError:
        st.error("Your API key was rejected. Use “Change API key” in the sidebar to enter a valid one.")
    except ImportError as e:
        st.error(f"A required package is missing: {e}. Run: python -m pip install -r requirements.txt")
    except Exception as e:  # noqa: BLE001
        st.error(f"Couldn't index this file: {e}")
    st.session_state.pop("vectorstore", None)
    st.session_state.pop("index_sig", None)
    return False


def suggestions_for(ftype: str) -> list:
    if ftype in ("csv", "excel"):
        return ["Give me an overview of this data",
                "What columns does this file have?",
                "What are the most important findings?"]
    return ["Summarize this document",
            "What are the key points?",
            "What topics does this document cover?"]


def render_chat(uid: str, cid: str, chat: dict, model_name: str, top_k: int, api_key: str) -> None:
    info = st.session_state.get("doc_info", {})
    label, icon, _ = TYPE_META.get(chat["type"], TYPE_META["pdf"])

    # ---- Header ----
    st.markdown(
        '<div class="app-header">'
        f'<div class="brand">{LOGO_SVG}<div><div class="title">{icon} {html.escape(chat["file"])}</div>'
        f'<div class="sub">{label} chat · answers grounded in this file</div></div></div>'
        '<div class="chip"><span class="dot"></span>API verified</div>'
        "</div>",
        unsafe_allow_html=True,
    )

    stats_slot = st.empty()

    def draw_stats() -> None:
        count = len(st.session_state.get("messages", []))
        stats_slot.markdown(
            '<div class="stats">'
            f'<div class="stat"><div class="s-label">Type</div><div class="s-value">{icon} {label}</div></div>'
            f'<div class="stat"><div class="s-label">{info["unit_name"]}</div><div class="s-value">{info["units"]:,}</div></div>'
            f'<div class="stat"><div class="s-label">Indexed chunks</div><div class="s-value">{info["chunks"]:,}</div></div>'
            f'<div class="stat"><div class="s-label">Messages</div><div class="s-value">{count:,}</div></div>'
            "</div>",
            unsafe_allow_html=True,
        )

    draw_stats()
    if info.get("note"):
        st.warning(info["note"])

    typed = st.chat_input("Ask a question about your file…")
    question = typed or st.session_state.pop("pending_q", None)

    messages = st.session_state.messages

    # ---- Suggestions on an empty chat ----
    if not messages and not question:
        st.markdown('<div class="sect-label">Try asking</div>', unsafe_allow_html=True)
        cols = st.columns(3, gap="small")
        for i, (col, text) in enumerate(zip(cols, suggestions_for(chat["type"]))):
            with col:
                st.button(text, key=f"sug_{i}", on_click=ask_cb, args=(text,), use_container_width=True)

    # ---- Toolbar ----
    if messages:
        st.download_button(
            "⬇  Download this chat (.md)",
            data=export_chat_md(chat["file"], messages),
            file_name=f'{chat["file"].rsplit(".", 1)[0]}_chat.md',
            mime="text/markdown",
            key=f"dl_{cid}_{len(messages)}",
        )

    # ---- Conversation ----
    for msg in messages:
        render_message(msg)

    if question:
        history_text = format_history(messages)
        user_msg = {"role": "user", "content": question, "time": _now()}
        messages.append(user_msg)
        render_message(user_msg)

        with st.chat_message("assistant", avatar="📄"):
            thinking = st.empty()
            thinking.markdown(THINKING_HTML, unsafe_allow_html=True)
            try:
                docs = retrieve(question, st.session_state.vectorstore, top_k)
                thinking.empty()
                answer = st.write_stream(
                    stream_answer(question, docs, history_text, model_name, api_key)
                )
                sources = docs_to_sources(docs)
                render_sources(sources)
                messages.append({"role": "assistant", "content": answer,
                                 "sources": sources, "time": _now()})
                save_messages(uid, cid, messages)
                draw_stats()
                st.rerun()  # refresh sidebar (title, counts) and show the download button
            except AuthenticationError:
                thinking.empty()
                st.error("Your API key was rejected. Use “Change API key” in the sidebar.")
            except RateLimitError:
                thinking.empty()
                st.error("Rate or quota limit reached. Check your OpenAI usage and try again shortly.")
            except Exception as e:  # noqa: BLE001
                thinking.empty()
                st.error(f"Something went wrong while answering: {e}")


def render_app() -> None:
    api_key = st.session_state.api_key
    uid = user_id(api_key)

    model_name, chunk_size, chunk_overlap, top_k = render_sidebar(uid, api_key)

    active = st.session_state.get("active_chat")
    chat = get_chat(uid, active) if active else None
    if active and chat is None:  # chat was deleted elsewhere
        reset_chat_state()
        active = None

    if chat is None:
        render_welcome(uid)
        return

    # load saved messages when a chat is opened
    if st.session_state.get("messages_for") != active:
        st.session_state.messages = list(chat.get("messages", []))
        st.session_state.messages_for = active

    if ensure_index(chat, active, api_key, chunk_size, chunk_overlap):
        render_chat(uid, active, chat, model_name, top_k, api_key)


# =============================================================================
# Router
# =============================================================================
if st.session_state.get("authenticated") and st.session_state.get("api_key"):
    render_app()
else:
    render_login()